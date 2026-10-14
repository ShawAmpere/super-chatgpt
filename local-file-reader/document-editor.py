"""Isolated document-library entry point. User code has current-user privileges."""
import importlib.metadata
import json
import os
from pathlib import Path
import subprocess
import sys
import traceback
import zipfile

PACKAGES = {"docx": "python-docx", "pptx": "python-pptx", "openpyxl": "openpyxl", "pypdf": "pypdf",
            "pymupdf": "PyMuPDF", "reportlab": "reportlab", "PIL": "Pillow", "odf": "odfpy"}
TEXT = {".csv", ".tsv", ".txt", ".md", ".html"}
ODF = {".odt": "application/vnd.oasis.opendocument.text", ".ods": "application/vnd.oasis.opendocument.spreadsheet",
       ".odp": "application/vnd.oasis.opendocument.presentation"}
FILTERS = {".pdf": "pdf", ".docx": "docx:Office Open XML Text", ".pptx": "pptx:Impress MS PowerPoint 2007 XML",
           ".xlsx": "xlsx:Calc MS Excel 2007 XML", ".doc": "doc:MS Word 97", ".xls": "xls:MS Excel 97",
           ".ppt": "ppt:MS PowerPoint 97", ".odt": "odt", ".ods": "ods", ".odp": "odp", ".rtf": "rtf", ".html": "html"}


def validate_document(filename):
    """Parse the staged output before publication; this is not visual layout QA."""
    filename = Path(filename)
    if filename.is_symlink() or not filename.is_file() or not 0 < filename.stat().st_size <= 64 * 1024 * 1024:
        raise ValueError("Document output must be a nonempty regular file of at most 64 MiB.")
    extension = filename.suffix.lower()
    if extension in {".docx", ".pptx", ".xlsx", *ODF}:
        with zipfile.ZipFile(filename) as archive:
            entries = archive.infolist()
            if len(entries) > 4096 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
                raise ValueError("Document output exceeds the 4096-entry/128 MiB expanded archive limit.")
    if extension == ".docx":
        from docx import Document
        document = Document(filename)
        return {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    if extension == ".pptx":
        from pptx import Presentation
        return {"slides": len(Presentation(filename).slides)}
    if extension == ".xlsx":
        from openpyxl import load_workbook
        document = load_workbook(filename, read_only=True)
        try:
            return {"sheets": document.sheetnames}
        finally:
            document.close()
    if extension == ".pdf":
        import pymupdf
        with pymupdf.open(filename) as document:
            if document.is_encrypted:
                raise ValueError("Encrypted PDF output is not supported.")
            return {"pages": len(document)}
    if extension in ODF:
        from odf.opendocument import load
        with zipfile.ZipFile(filename) as archive:
            if archive.read("mimetype").decode() != ODF[extension]:
                raise ValueError("OpenDocument content does not match its extension.")
        load(str(filename))
        return {"format": extension[1:]}
    if extension in {".doc", ".xls", ".ppt"}:
        with filename.open("rb") as stream:
            if stream.read(8) != bytes.fromhex("d0cf11e0a1b11ae1"):
                raise ValueError("Legacy Office output is not an OLE container.")
        return {"format": extension[1:], "validation": "container_signature_only"}
    if extension == ".rtf":
        if not filename.read_bytes().lstrip().startswith(b"{\\rtf"):
            raise ValueError("Output is not an RTF document.")
        return {"format": "rtf"}
    if extension in TEXT:
        text = filename.read_text(encoding="utf-8-sig")
        if "\x00" in text:
            raise ValueError("Text output contains NUL.")
        return {"characters": len(text)}
    raise ValueError("Unsupported document output extension.")


def run(job_path):
    job_path = Path(job_path)
    configuration = json.loads(job_path.read_text(encoding="utf-8"))
    work = job_path.parent
    converter_work = Path(configuration["converter_work_dir"])
    conversion_count = 0

    def convert_document(source, extension, destination=None):
        """Convert a stage file with a separate LibreOffice profile and output directory."""
        nonlocal conversion_count
        extension = "." + extension.lower().lstrip(".")
        if extension not in FILTERS:
            raise ValueError("Unsupported LibreOffice output format.")
        converter = os.environ.get("LOCAL_FILE_READER_LIBREOFFICE", r"C:\Program Files\LibreOffice\program\soffice.exe")
        if not Path(converter).is_file():
            raise RuntimeError("LibreOffice is not installed or LOCAL_FILE_READER_LIBREOFFICE is incorrect.")
        # The Windows console launcher waits for conversion and captures diagnostics.
        console_launcher = Path(converter).with_suffix(".com")
        if Path(converter).name.lower() == "soffice.exe" and console_launcher.is_file():
            converter = str(console_launcher)
        conversion_count += 1
        directory = converter_work / f"conversion-{conversion_count}"
        directory.mkdir()
        profile = (converter_work / f"profile-{conversion_count}").as_uri()
        result = subprocess.run([converter, f"-env:UserInstallation={profile}", "--headless", "--nologo", "--nodefault", "--convert-to",
                                 FILTERS[extension], "--outdir", str(directory), str(Path(source).resolve())],
                                capture_output=True, text=True, errors="replace", timeout=60,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        converted = directory / (Path(source).stem + extension)
        if result.returncode or not converted.is_file():
            raise RuntimeError(f"LibreOffice conversion failed (exit {result.returncode}): {result.stderr or result.stdout or 'no output produced'}")
        if destination:
            import shutil
            shutil.copyfile(converted, destination)
            return str(destination)
        return str(converted)

    scope = {"__name__": "__main__", "__file__": str(work / "user-code.py"), "INPUT_PATH": configuration.get("input_path"),
             "OUTPUT_PATH": configuration["output_path"], "WORK_DIR": str(work), "convert_document": convert_document}
    code = (work / "user-code.py").read_text(encoding="utf-8")
    exec(compile(code, scope["__file__"], "exec"), scope)
    validation = validate_document(configuration["output_path"])
    (work / "validated.json").write_text(json.dumps(validation, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    try:
        if sys.argv[1] == "capabilities":
            versions = {name: importlib.metadata.version(package) for name, package in PACKAGES.items()}
            for name in PACKAGES:
                __import__(name)
            print(json.dumps({"available": True, "libraries": versions}, ensure_ascii=False))
        else:
            run(sys.argv[1])
    except Exception:
        traceback.print_exc()
        sys.exit(1)
